"""Read the dev-package Excel file (one sheet per package: D01-D04) into
clean package dicts. Structure per sheet:
  row 0: package_id, title
  row 1: a provenance note (ignored)
  a 'Requirement ID' | 'Requirement' header, then requirement rows
  a blank row
  a 'Test ID' | 'Scenario' | 'Input' | 'Expected result' header, then test rows
"""

from openpyxl import load_workbook


def _section(rows, header_first_cell):
    """Find a header row by its first cell, yield dict rows until a blank row."""
    for i, row in enumerate(rows):
        if row and row[0] and str(row[0]).strip() == header_first_cell:
            headers = [str(h).strip() if h else "" for h in row]
            out = []
            for data in rows[i + 1:]:
                if not data or not data[0]:
                    break
                out.append(dict(zip(headers, data)))
            return out
    return []


def load_packages(file) -> list[dict]:
    wb = load_workbook(file, data_only=True)
    packages = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows or not rows[0][0]:
            continue
        package_id, title = rows[0][0], rows[0][1]

        reqs = _section(rows, "Requirement ID")
        tests = _section(rows, "Test ID")

        requirement_text = f"{package_id} — {title}\n\n" + "\n".join(
            f"{r['Requirement ID']}: {r['Requirement']}" for r in reqs
        )
        test_lines = [
            f"{t['Test ID']} | {t['Scenario']} — {t['Input']} | Expected: {t['Expected result']}"
            for t in tests
        ]
        packages.append({
            "name": f"{package_id} — {title}",
            "requirement_text": requirement_text,
            "test_cases_text": "\n".join(test_lines),
            "n_requirements": len(reqs),
            "n_tests": len(tests),
        })
    return packages