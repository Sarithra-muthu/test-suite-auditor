"""Export reviewer decisions from the database into one scoring workbook.

READ-ONLY: runs SELECT statements only; it never writes to the database.

Usage (from the test-suite-auditor folder, venv active):
    python -m scripts.export_for_scoring

The workbook is written OUTSIDE the repo (../scoring_exports/) so reviewer data
can never be committed by accident. Timestamps are UTC.

What goes in it
  Review progress    how far each reviewer has got on each package
  Agent-only output  every frozen AI finding / question, before any human decision
  Paired final       the human-confirmed set: accepted + edited AI findings + reviewer additions
  Rejected           AI findings the reviewer rejected, with their reason
  Clarifications     all clarification questions, with answers and "Assumption:" flags
  RTM                the frozen coverage matrix
  Decision history   every decision event (append-only audit trail)
  Warnings           anything that needs a human look before scoring
Human-only findings are not in the app, so they are not here (they come from the Excel worksheets).
"""

import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

EVAL_CODES = [f"P{n:02d}" for n in range(1, 11)]
TABLES = ("packages", "findings", "clarification_questions", "rtm_entries", "review_decisions")

SCORING_HEADERS = ["Matches reference issue ID", "Valid? (Yes/No/Duplicate)", "Scoring notes"]
LEGEND = ("Yellow columns are for you. 'Matches reference issue ID': the key's ID, or leave blank if it matches nothing. "
          "'Valid?': Yes = a real issue (in the key or adjudicated in); No = unsupported; Duplicate = repeats another row. "
          "Count each underlying issue once.")


# ---------------------------------------------------------------- helpers
def ts(v):
    return "" if v is None else str(v)[:19]


def code_of(name):
    parts = (name or "").split()
    return parts[0].strip() if parts else ""


def problems_text(raw):
    try:
        return "; ".join(json.loads(raw or "[]"))
    except Exception:
        return str(raw)


def yn(v):
    return "Yes" if v else "No"


def decision_label(decision, final_text, explanation):
    if decision == "rejected":
        return "rejected"
    if decision == "accepted":
        changed = bool(final_text and final_text.strip() and final_text.strip() != (explanation or "").strip())
        return "edited" if changed else "accepted"
    return decision or "pending"


def read_all(conn):
    cur = conn.cursor()
    out = {}
    for t in TABLES:
        cur.execute(f"SELECT * FROM {t}")
        out[t] = [dict(r) for r in cur.fetchall()]
    return out


def pick_frozen(packages, warnings):
    """One COMPLETE package per evaluation code; anything else is reported, not used."""
    by_code = {}
    for p in packages:
        code = code_of(p["name"])
        if code not in EVAL_CODES:
            continue
        status = (p.get("audit_status") or "complete").lower()
        if status != "complete":
            warnings.append(f"{code}: package #{p['id']} is {status.upper()} (excluded)")
            continue
        by_code.setdefault(code, []).append(p)
    chosen = {}
    for code in EVAL_CODES:
        rows = sorted(by_code.get(code, []), key=lambda r: r["id"])
        if not rows:
            warnings.append(f"{code}: no COMPLETE package found")
            continue
        if len(rows) > 1:
            warnings.append(f"{code}: {len(rows)} COMPLETE packages {[r['id'] for r in rows]}; "
                            f"using the first frozen, #{rows[0]['id']}")
        chosen[code] = rows[0]
    return chosen


# ---------------------------------------------------------------- transform
def build(tables, assignments):
    """Return ({sheet name: spec}, warnings). Pure Python, no database access."""
    warnings = []
    code_to_rev = {c: r for r, codes in assignments.items() for c in codes}
    frozen = pick_frozen(tables["packages"], warnings)
    pkg = {p["id"]: (code, p) for code, p in frozen.items()}
    order = {code: i for i, code in enumerate(EVAL_CODES)}

    def code(pid):
        return pkg[pid][0]

    findings = sorted((f for f in tables["findings"] if f["package_id"] in pkg),
                      key=lambda f: (order[code(f["package_id"])], f["id"]))
    questions = sorted((q for q in tables["clarification_questions"] if q["package_id"] in pkg),
                       key=lambda q: (order[code(q["package_id"])], q["id"]))
    rtm = sorted((r for r in tables["rtm_entries"] if r["package_id"] in pkg),
                 key=lambda r: (order[code(r["package_id"])], r["id"]))
    finding_pkg = {f["id"]: f["package_id"] for f in findings}

    history, latest, stray = [], {}, Counter()
    for d in sorted(tables["review_decisions"], key=lambda d: d["id"]):
        pid = finding_pkg.get(d["finding_id"])
        if pid is None:
            continue
        c = code(pid)
        assigned = code_to_rev.get(c)
        ok = d["reviewer_id"] == assigned
        history.append((d, c, assigned, ok))
        if ok:
            latest[d["finding_id"]] = d
        else:
            stray[(c, d["reviewer_id"], assigned)] += 1
    for (c, who, assigned), n in sorted(stray.items()):
        warnings.append(f"{c}: {n} decision(s) by '{who}' but the package is assigned to '{assigned}' "
                        f"(excluded from Paired final / Rejected)")

    ai = [f for f in findings if (f.get("source") or "ai") == "ai"]
    human = [f for f in findings if f.get("source") == "human"]
    ai_q = [q for q in questions if (q.get("source") or "ai") == "ai"]

    # ---- Review progress
    progress = []
    for c, p in frozen.items():
        pid = p["id"]
        mine = [f for f in ai if f["package_id"] == pid]
        labels = Counter()
        for f in mine:
            d = latest.get(f["id"])
            labels[decision_label(d["decision"], d["final_text"], f["explanation"]) if d else "pending"] += 1
        qs = [q for q in questions if q["package_id"] == pid]
        times = [d["created_at"] for d, cc, _, ok in history if cc == c and ok and d["created_at"]]
        state = "n/a (no AI findings)" if not mine else ("all decided" if not labels["pending"] else f"{labels['pending']} pending")
        progress.append([c, pid, code_to_rev.get(c, "(unassigned)"), len(mine), labels["accepted"], labels["edited"],
                         labels["rejected"], labels["pending"], state,
                         sum(1 for f in human if f["package_id"] == pid),
                         sum(1 for q in qs if (q.get("source") or "ai") == "ai"),
                         sum(1 for q in qs if q.get("source") == "human"),
                         sum(1 for q in qs if q.get("answer")), ts(max(times)) if times else ""])

    # ---- Agent-only output (before any human decision)
    agent = []
    for f in ai:
        agent.append([code(f["package_id"]), f["package_id"], "finding", f["id"], f["level"], f["category"],
                      f["test_id"] or "", f["source_quote"], f["explanation"], f["priority"],
                      yn(f["reference_valid"]), problems_text(f["validation_problems"]), "", "", ""])
    for q in ai_q:
        agent.append([code(q["package_id"]), q["package_id"], "clarification question", q["id"], "requirement",
                      "clarification", "", q["requirement_ref"], q["question"], "",
                      yn(q["reference_valid"]), problems_text(q["validation_problems"]), "", "", ""])

    # ---- Paired final (human-confirmed) and Rejected
    paired, rejected = [], []
    for f in ai:
        d = latest.get(f["id"])
        if not d:
            continue
        c = code(f["package_id"])
        lab = decision_label(d["decision"], d["final_text"], f["explanation"])
        if lab == "rejected":
            rejected.append([c, d["reviewer_id"], f["id"], f["category"], f["test_id"] or "", f["source_quote"],
                             f["explanation"], d["final_text"] or "", ts(d["created_at"])])
        elif lab in ("accepted", "edited"):
            final = d["final_text"] if lab == "edited" else f["explanation"]
            paired.append([c, d["reviewer_id"], f"AI finding, {lab}", f["id"], f["level"], f["category"],
                           f["test_id"] or "", f["source_quote"], final,
                           f["explanation"] if lab == "edited" else "", yn(f["reference_valid"]), "", "", ""])
    for f in human:
        c = code(f["package_id"])
        paired.append([c, code_to_rev.get(c, "(unassigned)"), "Reviewer addition", f["id"], f["level"], f["category"],
                       f["test_id"] or "", f["source_quote"], f["explanation"], "", "n/a", "", "", ""])
    paired.sort(key=lambda r: (order[r[0]], r[2], r[3]))

    # ---- Clarifications
    clar = []
    for q in questions:
        c = code(q["package_id"])
        ans = q.get("answer") or ""
        origin = "reviewer addition" if q.get("source") == "human" else "AI"
        clar.append([c, origin, code_to_rev.get(c, "(unassigned)"), q["question"], q["requirement_ref"],
                     q["status"], ans, yn(ans.strip().lower().startswith("assumption:")), "", "", ""])

    rtm_rows = [[code(r["package_id"]), r["requirement_rule"], ", ".join(json.loads(r["test_ids"] or "[]")),
                 r["coverage_status"]] for r in rtm]

    hist_rows = [[c, d["id"], d["finding_id"], d["reviewer_id"], assigned or "", yn(ok), d["decision"],
                  d["final_text"] or "", ts(d["created_at"])] for d, c, assigned, ok in history]

    warn_rows = [[w] for w in warnings] or [["None"]]

    sheets = {
        "Review progress": dict(
            note="Progress per frozen package. Pending = AI findings the assigned reviewer has not yet decided. Times are UTC.",
            headers=["Package", "Package id", "Assigned reviewer", "AI findings", "Accepted", "Edited", "Rejected",
                     "Pending", "State", "Reviewer-added findings", "AI questions", "Reviewer-added questions",
                     "Questions answered", "Last decision (UTC)"],
            rows=progress, widths=[9, 10, 16, 11, 10, 8, 10, 9, 22, 14, 11, 14, 12, 20]),
        "Agent-only output": dict(
            note="Frozen auditor output, before any human decision. This is the agent-only condition. " + LEGEND,
            headers=["Package", "Package id", "Type", "Row id", "Level", "Category", "Test ID", "Source quote / requirement ref",
                     "Explanation / question", "Priority", "Reference valid", "Validation problems"] + SCORING_HEADERS,
            rows=agent, widths=[9, 10, 18, 8, 12, 20, 12, 45, 60, 9, 11, 30, 18, 16, 28], scoring=3),
        "Paired final": dict(
            note="Human-confirmed findings only (accepted, edited, or added by the reviewer); pending and rejected AI findings are excluded. "
                 "Reviewer additions are attributed by package assignment. " + LEGEND,
            headers=["Package", "Reviewer", "Origin", "Row id", "Level", "Category", "Test ID", "Source quote",
                     "Final text", "Original AI text (if edited)", "Reference valid"] + SCORING_HEADERS,
            rows=paired, widths=[9, 10, 22, 8, 12, 20, 12, 45, 60, 45, 11, 18, 16, 28], scoring=3),
        "Rejected": dict(
            note="AI findings the reviewer rejected, with their stated reason (useful for precision and false-alarm analysis).",
            headers=["Package", "Reviewer", "Finding id", "Category", "Test ID", "Source quote", "AI explanation",
                     "Rejection reason", "Decided (UTC)"],
            rows=rejected, widths=[9, 10, 10, 20, 12, 45, 60, 40, 20]),
        "Clarifications": dict(
            note="All clarification questions. 'Assumption?' = the answer starts with 'Assumption:', meaning the reviewer's own guess, "
                 "not an owner-confirmed answer. " + LEGEND,
            headers=["Package", "Origin", "Reviewer", "Question", "Requirement ref", "Status", "Answer", "Assumption?"] + SCORING_HEADERS,
            rows=clar, widths=[9, 18, 10, 60, 16, 12, 45, 12, 18, 16, 28], scoring=3),
        "RTM": dict(
            note="Frozen coverage matrix. Linkage is model judgment and varies slightly between runs.",
            headers=["Package", "Requirement rule", "Linked tests", "Coverage"],
            rows=rtm_rows, widths=[9, 60, 45, 24]),
        "Decision history": dict(
            note="Every decision event, oldest first (append-only). 'Assigned?' = made by the reviewer the package is assigned to.",
            headers=["Package", "Decision id", "Finding id", "Reviewer", "Assigned reviewer", "Assigned?", "Decision",
                     "Final text / reason", "Time (UTC)"],
            rows=hist_rows, widths=[9, 11, 10, 11, 16, 10, 11, 60, 20]),
        "Warnings": dict(
            note="Check these before scoring.", headers=["Warning"], rows=warn_rows, widths=[120]),
    }
    return sheets, warnings


# ---------------------------------------------------------------- write
def write_xlsx(path, sheets):
    f_note = Font(name="Arial", size=9, italic=True, color="595959")
    f_head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    f_body = Font(name="Arial", size=10)
    fill_head = PatternFill("solid", start_color="305496")
    fill_score = PatternFill("solid", start_color="FFF2CC")
    side = Side(style="thin", color="BFBFBF")
    box = Border(left=side, right=side, top=side, bottom=side)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    wb.remove(wb.active)
    for title, s in sheets.items():
        ws = wb.create_sheet(title)
        headers, rows, n_score = s["headers"], s["rows"], s.get("scoring", 0)
        ws["A1"] = s["note"]
        ws["A1"].font = f_note
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(2, c, h)
            cell.font, cell.fill, cell.border = f_head, fill_head, box
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = s["widths"][c - 1]
        for r, row in enumerate(rows, start=3):
            for c, v in enumerate(row, start=1):
                cell = ws.cell(r, c, v)
                cell.font, cell.alignment, cell.border = f_body, wrap, box
                if n_score and c > len(headers) - n_score:
                    cell.fill = fill_score
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{max(2, len(rows) + 2)}"
        if n_score and rows:
            col = get_column_letter(len(headers) - 1)
            dv = DataValidation(type="list", formula1='"Yes,No,Duplicate"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}3:{col}{len(rows) + 2}")
    wb.save(path)


def main():
    from app.auth import PACKAGE_ASSIGNMENTS
    from app.db import get_connection

    conn = get_connection()
    try:
        tables = read_all(conn)
    finally:
        conn.close()
    sheets, warnings = build(tables, PACKAGE_ASSIGNMENTS)

    out_dir = Path(__file__).resolve().parent.parent.parent / "scoring_exports"
    out_dir.mkdir(exist_ok=True)
    path = out_dir / f"scoring_export_{datetime.now(timezone.utc):%Y%m%d_%H%M}.xlsx"
    write_xlsx(path, sheets)

    print("\nReview progress")
    for r in sheets["Review progress"]["rows"]:
        print(f"  {r[0]}  {r[2]:<10} AI findings {r[3]:>2}  accepted {r[4]}  edited {r[5]}  rejected {r[6]}  -> {r[8]}")
    print("\nWarnings:", "none" if not warnings else "")
    for w in warnings:
        print("  -", w)
    print("\nSaved:", path)


if __name__ == "__main__":
    main()
